import io
import fitz
import pytest
from PIL import Image
import openpyxl
from pathlib import Path

from app.services.pdf_service import pdf_service
from app.services.image_service import image_service
from app.services.spreadsheet_service import spreadsheet_service
from app.services.text_service import text_service


@pytest.fixture
def sample_pdf(tmp_path) -> Path:
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Antigravity Document Converter Test Document\nHello World!", fontsize=14)
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def sample_image(tmp_path) -> Path:
    img_path = tmp_path / "sample.png"
    img = Image.new("RGBA", (100, 100), (255, 0, 0, 255))
    img.save(str(img_path))
    return img_path


@pytest.fixture
def sample_csv(tmp_path) -> Path:
    csv_path = tmp_path / "sample.csv"
    csv_path.write_text("Name,Age,Role\nAlice,30,Developer\nBob,25,Designer\nCharlie,35,Manager\n", encoding="utf-8")
    return csv_path


@pytest.fixture
def sample_markdown(tmp_path) -> Path:
    md_path = tmp_path / "sample.md"
    md_path.write_text("# Test Title\n\nThis is a **markdown** paragraph.\n\n- Item 1\n- Item 2\n", encoding="utf-8")
    return md_path


def test_pdf_to_txt(sample_pdf, tmp_path):
    out_dir = tmp_path / "out_txt"
    out_dir.mkdir()
    res = pdf_service.convert(sample_pdf, "txt", out_dir)
    assert len(res) == 1
    content = res[0].read_text(encoding="utf-8")
    assert "Antigravity Document Converter" in content
    assert "Hello World!" in content


def test_pdf_to_docx(sample_pdf, tmp_path):
    out_dir = tmp_path / "out_docx"
    out_dir.mkdir()
    res = pdf_service.convert(sample_pdf, "docx", out_dir)
    assert len(res) == 1
    assert res[0].exists()
    assert res[0].stat().st_size > 0


def test_pdf_to_png(sample_pdf, tmp_path):
    out_dir = tmp_path / "out_png"
    out_dir.mkdir()
    res = pdf_service.convert(sample_pdf, "png", out_dir)
    assert len(res) == 1
    assert res[0].name.endswith(".png")
    assert res[0].stat().st_size > 0


def test_pdf_watermark(sample_pdf, tmp_path):
    out_dir = tmp_path / "out_watermark"
    out_dir.mkdir()
    watermarked = pdf_service.add_watermark(sample_pdf, out_dir, watermark_text="PRIVATE", opacity=0.5)
    assert watermarked.exists()
    assert watermarked.stat().st_size > 0


def test_pdf_compress_options(sample_pdf, tmp_path):
    out_dir = tmp_path / "out_compress"
    out_dir.mkdir()
    compressed = pdf_service.compress_pdf(
        sample_pdf,
        out_dir,
        quality=50,
        dpi=96,
        grayscale=True,
        strip_metadata=True
    )
    assert compressed.exists()
    assert compressed.stat().st_size > 0


def test_pdf_add_page_numbers(sample_pdf, tmp_path):
    out_dir = tmp_path / "out_num"
    out_dir.mkdir()
    numbered = pdf_service.add_page_numbers(sample_pdf, out_dir, format_pattern="Page {page} of {total}")
    assert numbered.exists()
    assert numbered.stat().st_size > 0


def test_pdf_delete_pages(sample_pdf, tmp_path):
    out_dir = tmp_path / "out_del"
    out_dir.mkdir()
    # sample_pdf has at least 1 page, let's test with a 2-page pdf
    import fitz
    doc = fitz.open(str(sample_pdf))
    doc.new_page()
    multi_p = tmp_path / "multi.pdf"
    doc.save(str(multi_p))
    doc.close()

    result = pdf_service.delete_pages(multi_p, out_dir, pages_to_delete=[1])
    assert result.exists()
    doc2 = fitz.open(str(result))
    assert len(doc2) == 1
    doc2.close()



def test_image_png_to_jpg(sample_image, tmp_path):
    out_dir = tmp_path / "out_img"
    out_dir.mkdir()
    res = image_service.convert(sample_image, "jpg", out_dir, options={"quality": 85})
    assert len(res) == 1
    assert res[0].name.endswith(".jpg")
    with Image.open(str(res[0])) as img:
        assert img.format == "JPEG"


def test_image_resize(sample_image, tmp_path):
    out_dir = tmp_path / "out_resize"
    out_dir.mkdir()
    res = image_service.convert(sample_image, "png", out_dir, options={"width": 50, "height": 50})
    assert len(res) == 1
    with Image.open(str(res[0])) as img:
        assert img.size == (50, 50)


def test_spreadsheet_csv_to_xlsx(sample_csv, tmp_path):
    out_dir = tmp_path / "out_sheet"
    out_dir.mkdir()
    res = spreadsheet_service.convert(sample_csv, "xlsx", out_dir)
    assert len(res) == 1
    wb = openpyxl.load_workbook(str(res[0]))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Name"
    assert ws.cell(row=2, column=1).value == "Alice"


def test_spreadsheet_preview(sample_csv):
    preview = spreadsheet_service.preview_data(sample_csv)
    assert preview["total_rows"] == 3
    assert "Name" in preview["columns"]
    assert preview["rows"][0]["Name"] == "Alice"


def test_markdown_to_html(sample_markdown, tmp_path):
    out_dir = tmp_path / "out_md"
    out_dir.mkdir()
    res = text_service.convert(sample_markdown, "html", out_dir)
    assert len(res) == 1
    html = res[0].read_text(encoding="utf-8")
    assert "Test Title</h1>" in html
    assert "<strong>markdown</strong>" in html
