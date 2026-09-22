from pathlib import Path
from typing import List, Dict, Any, Optional
import markdown
from bs4 import BeautifulSoup
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from xhtml2pdf import pisa

from app.services.base_converter import BaseConverter
from app.services.office_service import office_service


class TextService(BaseConverter):
    @property
    def supported_inputs(self) -> List[str]:
        return ["txt", "md", "html", "htm", "rtf"]

    @property
    def supported_outputs(self) -> Dict[str, List[str]]:
        return {
            "txt": ["pdf", "docx"],
            "md": ["html", "pdf", "docx"],
            "html": ["pdf", "docx", "txt"],
            "htm": ["pdf", "docx", "txt"],
            "rtf": ["docx", "txt", "pdf"]
        }

    def convert(
        self,
        input_path: Path,
        target_format: str,
        output_dir: Path,
        options: Optional[Dict[str, Any]] = None
    ) -> List[Path]:
        ext = input_path.suffix.lower().lstrip(".")
        target = target_format.lower().lstrip(".")
        stem = input_path.stem
        out_path = output_dir / f"{stem}.{target}"

        # Read text content
        raw_text = input_path.read_text(encoding="utf-8", errors="replace")

        # Markdown conversions
        if ext == "md":
            if target == "html":
                html_body = markdown.markdown(raw_text, extensions=["extra", "tables", "fenced_code", "toc"])
                full_html = (
                    "<!DOCTYPE html>\n<html><head><meta charset='utf-8'>"
                    f"<title>{stem}</title>"
                    "<style>body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; margin: 2rem auto; max-width: 800px; padding: 0 1rem; line-height: 1.6; color: #1e293b; }</style>"
                    f"</head><body>{html_body}</body></html>"
                )
                with open(out_path, "w", encoding="utf-8") as f:
                    f.write(full_html)
                return [out_path]

            elif target == "pdf":
                html_body = markdown.markdown(raw_text, extensions=["extra", "tables", "fenced_code"])
                full_html = f"<html><head><style>body {{ font-family: Helvetica, sans-serif; font-size: 11pt; line-height: 1.5; color: #1e293b; }}</style></head><body>{html_body}</body></html>"
                with open(out_path, "wb") as f:
                    pisa_status = pisa.CreatePDF(html_body, dest=f)
                if pisa_status.err:
                    # Fallback to reportlab preformatted text
                    return self._txt_to_pdf(raw_text, stem, out_path)
                return [out_path]

            elif target == "docx":
                return self._markdown_to_docx(raw_text, out_path)

        # Plain text conversions
        elif ext == "txt":
            if target == "pdf":
                return self._txt_to_pdf(raw_text, stem, out_path)
            elif target == "docx":
                doc = docx.Document()
                for line in raw_text.splitlines():
                    doc.add_paragraph(line)
                doc.save(str(out_path))
                return [out_path]

        # HTML conversions
        elif ext in ("html", "htm"):
            if target == "pdf":
                with open(out_path, "wb") as f:
                    pisa_status = pisa.CreatePDF(raw_text, dest=f)
                if pisa_status.err:
                    soup = BeautifulSoup(raw_text, "html.parser")
                    return self._txt_to_pdf(soup.get_text(), stem, out_path)
                return [out_path]

            elif target == "txt":
                soup = BeautifulSoup(raw_text, "html.parser")
                clean_text = soup.get_text(separator="\n\n")
                with open(out_path, "w", encoding="utf-8") as f:
                    f.write(clean_text)
                return [out_path]

            elif target == "docx":
                soup = BeautifulSoup(raw_text, "html.parser")
                doc = docx.Document()
                for elem in soup.find_all(['h1', 'h2', 'h3', 'p', 'li']):
                    if elem.name == 'h1':
                        doc.add_heading(elem.get_text(), level=1)
                    elif elem.name == 'h2':
                        doc.add_heading(elem.get_text(), level=2)
                    elif elem.name == 'h3':
                        doc.add_heading(elem.get_text(), level=3)
                    elif elem.name == 'li':
                        doc.add_paragraph(elem.get_text(), style='List Bullet')
                    else:
                        doc.add_paragraph(elem.get_text())
                doc.save(str(out_path))
                return [out_path]

        # RTF conversions
        elif ext == "rtf":
            if office_service.has_libreoffice:
                return office_service.convert_via_libreoffice(input_path, target, output_dir)
            else:
                # Basic RTF to text extraction
                text = self._strip_rtf(raw_text)
                if target == "txt":
                    with open(out_path, "w", encoding="utf-8") as f:
                        f.write(text)
                    return [out_path]
                elif target == "docx":
                    doc = docx.Document()
                    for line in text.splitlines():
                        if line.strip():
                            doc.add_paragraph(line)
                    doc.save(str(out_path))
                    return [out_path]

        raise ValueError(f"Unsupported conversion: {ext} -> {target}")

    def _txt_to_pdf(self, text: str, title: str, out_path: Path) -> List[Path]:
        doc = SimpleDocTemplate(str(out_path), pagesize=letter, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
        styles = getSampleStyleSheet()
        elements = []
        
        body_style = ParagraphStyle(
            'TxtBody',
            parent=styles['Normal'],
            fontName='Courier',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#1e293b")
        )
        
        # Split into blocks to avoid layout overflows
        lines = text.splitlines()
        chunk_size = 50
        for i in range(0, len(lines), chunk_size):
            chunk = "\n".join(lines[i:i + chunk_size])
            elements.append(Preformatted(chunk, body_style))
            elements.append(Spacer(1, 5))

        doc.build(elements)
        return [out_path]

    def _markdown_to_docx(self, text: str, out_path: Path) -> List[Path]:
        doc = docx.Document()
        for line in text.splitlines():
            sline = line.strip()
            if sline.startswith("# "):
                doc.add_heading(sline[2:], level=1)
            elif sline.startswith("## "):
                doc.add_heading(sline[3:], level=2)
            elif sline.startswith("### "):
                doc.add_heading(sline[4:], level=3)
            elif sline.startswith("- ") or sline.startswith("* "):
                doc.add_paragraph(sline[2:], style='List Bullet')
            elif sline.startswith("> "):
                doc.add_paragraph(sline[2:], style='Quote')
            elif sline:
                doc.add_paragraph(line)
        doc.save(str(out_path))
        return [out_path]

    def _strip_rtf(self, rtf_text: str) -> str:
        """Lightweight RTF command stripper when LibreOffice is not installed."""
        import re
        # Remove groups and rtf tokens
        text = re.sub(r'{\*?\\[^{}]+}|[{}]|\\\n?[A-Za-z]+ ?|\\[\'"][0-9a-fA-F]{2}', '', rtf_text)
        return text.strip()


text_service = TextService()
